"""Genera el MP Cyber Monday 2026 B2B Argentina de Pickit (dos escenarios) con estética Zenda.

Uso: python3 build_mp_cyber_b2b_2026.py <salida.xlsx>
Brief: briefs/pickit-cyber-monday-2026-b2b.md

Reglas de esta planilla (pedido de Ariana):
- La proyección va limpia; los benchmarks van en una tabla de referencia aparte, más abajo.
- El único texto visible son leyendas para el cliente (fuente de los datos y aclaraciones).
- Meta es UNA campaña con dos grupos de anuncios.
- La inversión se compara contra un día común de B2B Argentina (real de plataformas vía SeePM).
- BGT Diario mantiene el control de consumo que armó Milly.
"""
import os, sys, datetime as dt
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, ColorScaleRule
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from zenda_xlsx import *  # noqa: F401,F403

OUT = sys.argv[1]
wb = Workbook()

# ---------------------------------------------------------------------
# Datos reales (SeePM, 15/09 al 14/10/2026) — base de benchmarks y día común
# ---------------------------------------------------------------------
PERIODO = "15/09 al 14/10/2026"
REAL = {  # gasto, impresiones, clics, conversiones, alcance
    "LinkedIn": dict(ref="Boost de Contenido AR & MX (imagen y video)", gasto=394682.10, imp=37552, clics=142, conv=None),
    "Meta": dict(ref="AR_Hunting&SameDay_H2_FB_Conversiones_2026", gasto=259305.43, imp=27970, clics=320, conv=5),
    "Google": dict(ref="Search B2B AR (Flex/SameDay, E-commerce, Fintech)", gasto=2984777.29, imp=23805, clics=1365, conv=76),
}
DIAS_PERIODO = 30

# =====================================================================
# SUPUESTOS
# =====================================================================
S = wb.active
S.title = "Supuestos"
sheet(S, {"A": 2.6, "B": 40, "C": 18, "D": 18, "E": 16, "F": 16, "G": 14, "H": 14, "I": 14, "J": 14, "K": 14, "L": 14})
brand(S, "Supuestos del Media Plan")

band(S, 6, 2, 5, "FECHAS")
header(S, 7, 2, ["Etapa", "Inicio", "Fin", "Días"])
fechas = [("Previa al evento", dt.date(2026, 10, 26), dt.date(2026, 11, 1)),
          ("Cyber Monday", dt.date(2026, 11, 2), dt.date(2026, 11, 4))]
for i, (n, a, b) in enumerate(fechas):
    r = 8 + i
    cell(S, f"B{r}", n, bold=True)
    inp(S, f"C{r}", a, DATE); inp(S, f"D{r}", b, DATE)
    cell(S, f"E{r}", f"=D{r}-C{r}+1", INT)
total_row(S, 10, 2, 5)
put(S, "B10", "Flight total", F(bold=True), bg=L_VERDE, box=True)
put(S, "C10", "=C8", F(bold=True), DATE, L_VERDE, box=True)
put(S, "D10", "=D9", F(bold=True), DATE, L_VERDE, box=True)
put(S, "E10", "=E8+E9", F(bold=True), INT, L_VERDE, box=True)
INI, FIN, DIAS = "Supuestos!$C$10", "Supuestos!$D$10", "Supuestos!$E$10"

band(S, 12, 2, 6, "DÍA COMÚN B2B ARGENTINA")
header(S, 13, 2, ["Medio", "Inversión del período", "Días", "Inversión diaria", ""])
for i, (med, k) in enumerate((("LinkedIn Ads", "LinkedIn"), ("Meta Ads", "Meta"), ("Google Ads", "Google"))):
    r = 14 + i
    bg = ZEBRA if i % 2 else None
    cell(S, f"B{r}", med, bg=bg, bold=True)
    inp(S, f"C{r}", REAL[k]["gasto"], MONEY)
    inp(S, f"D{r}", DIAS_PERIODO, INT)
    cell(S, f"E{r}", f"=IFERROR(C{r}/D{r},\"\")", MONEY, bg)
    cell(S, f"F{r}", None, bg=bg)
total_row(S, 17, 2, 6)
put(S, "B17", "Día común B2B Argentina", F(bold=True), bg=L_VERDE, box=True)
put(S, "C17", "=SUM(C14:C16)", F(bold=True), MONEY, L_VERDE, box=True)
put(S, "E17", "=SUM(E14:E16)", F(bold=True, color=VERDE_INK), MONEY, L_VERDE, box=True)
DIA_COMUN = "Supuestos!$E$17"
note(S, "B18", f"Fuente: inversión real de las campañas B2B de Argentina en Google Ads, Meta Ads y LinkedIn Ads, {PERIODO}. LinkedIn incluye el boost de contenido AR & MX.")

band(S, 20, 2, 6, "PRESUPUESTO POR ESCENARIO")
header(S, 21, 2, ["Escenario", "Inversión diaria vs. día común", "Inversión diaria", "Inversión total", ""])
PRES, MULT = {}, {}
for i, (n, m) in enumerate((("Conservador", 1.5), ("Arriesgado", 3.0))):
    r = 22 + i
    bg = ZEBRA if i % 2 else None
    cell(S, f"B{r}", n, bold=True, bg=bg)
    inp(S, f"C{r}", m, '0.0"x"')
    cell(S, f"D{r}", f"={DIA_COMUN}*C{r}", MONEY, bg)
    cell(S, f"E{r}", f"=ROUND(D{r}*{DIAS},-3)", MONEY, bg, bold=True)
    cell(S, f"F{r}", None, bg=bg)
    PRES[n], MULT[n] = f"Supuestos!$E${r}", f"Supuestos!$C${r}"
note(S, "B24", "El presupuesto total sale de multiplicar la inversión de un día común por el factor del escenario y por los días del flight.")

band(S, 26, 2, 5, "AUDIENCIA DE LINKEDIN")
cell(S, "B27", "Tamaño de la audiencia (miembros)", bold=True)
inp(S, "D27", 120000, INT)
cell(S, "B28", "Cobertura máxima en el flight", bold=True)
inp(S, "D28", 0.5, PCT0)
cell(S, "B29", "Alcance máximo", bold=True)
cell(S, "D29", "=D27*D28", INT, bold=True)
cell(S, "B30", "Frecuencia mínima por persona", bold=True)
inp(S, "D30", 2.5, DEC1)
AUD_ALC, FREC_MIN = "Supuestos!$D$29", "Supuestos!$D$30"
note(S, "B31", "Alcance = impresiones / frecuencia mínima, hasta el alcance máximo de la audiencia.")

band(S, 33, 2, 5, "AJUSTE DE COSTOS POR TEMPORADA")
cell(S, "B34", "Suba de CPM y CPC durante el flight", bold=True)
inp(S, "D34", 0.15, PCT0)
AJ = "Supuestos!$D$34"
note(S, "B35", "Las subastas se encarecen en la semana del Cyber Monday porque todos los anunciantes aumentan su inversión.")

band(S, 37, 2, 12, "DISTRIBUCIÓN DIARIA DE LA INVERSIÓN")
header(S, 38, 2, [""] + [f"Día {i+1}" for i in range(10)])
cell(S, "B39", "Fecha", bold=True)
cell(S, "B40", "% de la inversión", bold=True)
pesos = [0.07, 0.07, 0.08, 0.08, 0.08, 0.08, 0.09, 0.15, 0.15, 0.15]
for i, w in enumerate(pesos):
    col = L(3 + i)
    cell(S, f"{col}39", f"=Supuestos!$C$8+{i}", "ddd dd/mm")
    inp(S, f"{col}40", w, PCT0)
total_row(S, 41, 2, 12)
put(S, "B41", "Total", F(bold=True), bg=L_VERDE, box=True)
put(S, "C41", "=SUM(C40:L40)", F(bold=True), PCT0, L_VERDE, box=True)
note(S, "B42", "La inversión crece en la previa y los tres días del Cyber Monday concentran el 45 %.")
CURVA_FECHA, CURVA_PCT = 39, 40

band(S, 44, 2, 5, "ESCENARIO A CONTROLAR EN EL BGT")
cell(S, "B45", "Escenario aprobado", bold=True)
x = inp(S, "D45", "Conservador", None)
x.alignment = Alignment(horizontal="center", vertical="center")
dv = DataValidation(type="list", formula1='"Conservador,Arriesgado"', allow_blank=False)
S.add_data_validation(dv); dv.add("D45")
ESC = "Supuestos!$D$45"

# =====================================================================
# MP
# =====================================================================
M = wb.create_sheet("MP Cyber B2B", 0)
widths = {"A": 2.6, "B": 38, "C": 26, "D": 12, "E": 13, "F": 28, "G": 7, "H": 13, "I": 12, "J": 12,
          "K": 13, "L": 9, "M": 10, "N": 10, "O": 12, "P": 12}
sheet(M, widths)
brand(M, "pickit  |  Media Plan Cyber Monday B2B · Argentina 2026")

COLS = ["Campaña", "Grupo de anuncios", "Medio", "Objetivo", "Audiencia", "%", "Inversión", "Inv. diaria",
        "Impresiones", "Clics", "CTR", "Alcance", "Frecuencia", "Conversiones", "CPA"]
NC = len(COLS)            # 15 columnas: B..P
LASTC = 1 + NC            # 16 = P

# filas de la tabla de benchmarks (se fija abajo, se referencia en la proyección)
BENCH_ROW = {"LinkedIn": None, "Meta": None, "Google": None}

DET = 29
band(M, DET, 2, LASTC, "PROYECCIÓN POR CAMPAÑA", h=26)
row = DET + 2

# estructura: (campaña, grupo, medio, objetivo, audiencia, %cons, %arr, bench, linkedin?)
campanias = [
    ("AR_B2B_LinkedIn_Boost_Awareness_Cyber2026", "Boost de posts", "LinkedIn Ads", "Reconocimiento",
     "Audiencia guardada B2B", 0.30, 0.20, "LinkedIn", True),
    ("AR_B2B_FB_Trafico_Cyber2026", None, "Meta Ads", "Tráfico web", None, None, None, "Meta", False),  # cabecera de campaña
    (None, "Público similar 1-3 % (BBDD)", "Meta Ads", "Tráfico web", "Similar a la base de datos", 0.20, 0.30, "Meta", False),
    (None, "Remarketing BBDD", "Meta Ads", "Tráfico web", "Base de datos de clientes y leads", 0.10, 0.10, "Meta", False),
    ("AR_B2B_GL_PMax_AliadoLogistico_Cyber2026", "Grupo de recursos único", "Google Ads", "Conversiones",
     "Señales: términos de búsqueda + remarketing BBDD", 0.40, 0.40, "Google", False),
]

TOT, ROWS, LI, META_HDR = {}, {}, {}, {}
plan_rows = {}
bench_cells = {}  # se completan después con las celdas reales


def bref(k, col):
    return f"${col}${BENCH_ROW[k]}"


# Primero calculamos dónde va la tabla de benchmarks: después de los dos escenarios
rows_per_esc = 3 + len(campanias) + 1 + 2      # banda + grupos + header + filas + total + blanco
BENCH_START = row + 2 * rows_per_esc + 1
BENCH_ROW["LinkedIn"], BENCH_ROW["Meta"], BENCH_ROW["Google"] = BENCH_START + 2, BENCH_START + 3, BENCH_START + 4

GRUPOS = [(2, 6, "CAMPAÑA", NEGRO, "FFFFFF"), (7, 9, "INVERSIÓN", VERDE_INK, "FFFFFF"),
          (10, 16, "RESULTADOS ESTIMADOS", ROSA, "FFFFFF")]

for escn, idx in (("Conservador", 5), ("Arriesgado", 6)):
    band(M, row, 2, LASTC, f"Escenario {escn.lower()}", 11, CYAN, NEGRO, h=22); row += 1
    groups(M, row, GRUPOS); row += 1
    header(M, row, 2, COLS); row += 1
    r0 = row
    leaf = []
    k_zebra = 0
    for cp in campanias:
        camp, grupo, med, obj, aud, pc, pa, bk, li = cp
        r = row
        if grupo is None:  # cabecera de campaña de Meta: suma sus grupos
            META_HDR[escn] = r
            for c in range(2, LASTC + 1):
                M.cell(r, c).fill = fill(L_CYAN); M.cell(r, c).border = BOX
            put(M, f"B{r}", camp, F(bold=True), bg=L_CYAN, box=True)
            put(M, f"C{r}", "Total campaña (2 grupos)", F(bold=True, color=MUTED), bg=L_CYAN, box=True)
            put(M, f"D{r}", med, F(bold=True), bg=L_CYAN, box=True)
            put(M, f"E{r}", obj, F(bold=True), bg=L_CYAN, box=True)
            for col, fm in (("G", PCT0), ("H", MONEY), ("I", MONEY), ("J", INT), ("K", INT), ("O", DEC1)):
                put(M, f"{col}{r}", f"=SUM({col}{r+1}:{col}{r+2})", F(bold=True), fm, L_CYAN, box=True)
            put(M, f"L{r}", f"=IFERROR(K{r}/J{r},\"\")", F(bold=True), PCT2, L_CYAN, box=True)
            put(M, f"P{r}", f"=IFERROR(H{r}/O{r},\"\")", F(bold=True), MONEY, L_CYAN, box=True)
            for col in "MN":
                put(M, f"{col}{r}", "–", F(bold=True, color=MUTED), bg=L_CYAN, al="right", box=True)
            row += 1
            continue
        bg = ZEBRA if k_zebra % 2 else None
        k_zebra += 1
        sub = camp is None
        cell(M, f"B{r}", camp if camp else "", bg=bg, bold=True)
        cell(M, f"C{r}", ("   ↳ " if sub else "") + grupo, bg=bg)
        cell(M, f"D{r}", med, bg=bg); cell(M, f"E{r}", obj, bg=bg); cell(M, f"F{r}", aud, bg=bg, color=MUTED)
        inp(M, f"G{r}", cp[idx], PCT0)
        cell(M, f"H{r}", f"=G{r}*{PRES[escn]}", MONEY, bg, bold=True)
        cell(M, f"I{r}", f"=IFERROR(H{r}/{DIAS},\"\")", MONEY, bg)
        cpm, cpc, cr = bref(bk, "H"), bref(bk, "I"), bref(bk, "K")
        cell(M, f"J{r}", f"=IFERROR(H{r}/({cpm}*(1+{AJ}))*1000,\"\")", INT, bg)
        cell(M, f"K{r}", f"=IFERROR(H{r}/({cpc}*(1+{AJ})),\"\")", INT, bg, bold=True)
        cell(M, f"L{r}", f"=IFERROR(K{r}/J{r},\"\")", PCT2, bg)
        if li:
            cell(M, f"M{r}", f"=MIN({AUD_ALC},J{r}/{FREC_MIN})", INT, bg, bold=True)
            cell(M, f"N{r}", f"=IFERROR(J{r}/M{r},\"\")", DEC1, bg, bold=True)
            cell(M, f"O{r}", "–", bg=bg, al="right", color=MUTED)
            cell(M, f"P{r}", "–", bg=bg, al="right", color=MUTED)
            LI[escn] = r
        else:
            cell(M, f"M{r}", "–", bg=bg, al="right", color=MUTED)
            cell(M, f"N{r}", "–", bg=bg, al="right", color=MUTED)
            cell(M, f"O{r}", f"=IFERROR(K{r}*{cr},\"\")", DEC1, bg)
            cell(M, f"P{r}", f"=IFERROR(H{r}/O{r},\"\")", MONEY, bg)
        leaf.append(r)
        row += 1
    r1 = row - 1
    T = row
    total_row(M, T, 2, LASTC)
    put(M, f"B{T}", f"Total escenario {escn.lower()}", F(bold=True), bg=L_VERDE, box=True)
    rng = ",".join(f"{{c}}{x}" for x in leaf)
    for col, fm in (("G", PCT0), ("H", MONEY), ("I", MONEY), ("J", INT), ("K", INT), ("O", DEC1)):
        put(M, f"{col}{T}", "=SUM(" + rng.format(c=col).replace("{c}", col) + ")",
            F(bold=True, color=VERDE_INK if col == "H" else NEGRO), fm, L_VERDE, box=True)
    put(M, f"L{T}", f"=IFERROR(K{T}/J{T},\"\")", F(bold=True), PCT2, L_VERDE, box=True)
    nonli = ",".join(f"H{x}" for x in leaf if x != LI[escn])
    put(M, f"P{T}", f"=IFERROR(SUM({nonli})/O{T},\"\")", F(bold=True), MONEY, L_VERDE, box=True)
    TOT[escn], ROWS[escn] = T, leaf
    plan_rows[escn] = leaf
    row += 2
note(M, f"B{row-1}", "Proyección con los benchmarks de la tabla de referencia y una suba de costos de 15 % por la temporada. El objetivo es presencia: alcance, frecuencia y clics; las conversiones son un indicador secundario.")

# ---- Benchmarks de referencia (abajo, separados de la proyección)
assert row <= BENCH_START, (row, BENCH_START)
b0 = BENCH_START
band(M, b0, 2, 14, "BENCHMARKS DE REFERENCIA", h=24)
header(M, b0 + 1, 2, ["Medio", "Campaña de referencia", None, None, "Período", None, "CPM", "CPC", "CTR", "CR", "Se usa en", None, None])
M.merge_cells(start_row=b0 + 1, start_column=3, end_row=b0 + 1, end_column=5)
M.merge_cells(start_row=b0 + 1, start_column=6, end_row=b0 + 1, end_column=7)
M.merge_cells(start_row=b0 + 1, start_column=12, end_row=b0 + 1, end_column=14)
uso = {"LinkedIn": "Boost de posts", "Meta": "Público similar y remarketing", "Google": "Performance Max"}
for i, k in enumerate(("LinkedIn", "Meta", "Google")):
    r = BENCH_ROW[k]
    d = REAL[k]
    bg = ZEBRA if i % 2 else None
    cell(M, f"B{r}", {"LinkedIn": "LinkedIn Ads", "Meta": "Meta Ads", "Google": "Google Ads"}[k], bg=bg, bold=True)
    for col in "CDE": cell(M, f"{col}{r}", None, bg=bg)
    M.merge_cells(f"C{r}:E{r}")
    cell(M, f"C{r}", d["ref"], bg=bg)
    for col in "FG": cell(M, f"{col}{r}", None, bg=bg)
    M.merge_cells(f"F{r}:G{r}")
    cell(M, f"F{r}", PERIODO, bg=bg, color=MUTED)
    cpm = round(d["gasto"] / d["imp"] * 1000, 2)
    cpc = round(d["gasto"] / d["clics"], 2)
    if k == "Google":  # PMax nueva: CPC y CR de Search; CTR estimado de PMax
        cpm = round(cpc * 0.01 * 1000, 2)
    inp(M, f"H{r}", cpm, MONEY); inp(M, f"I{r}", cpc, MONEY)
    cell(M, f"J{r}", f"=IFERROR(H{r}/I{r}/1000,\"\")", PCT2, bg)
    if d["conv"]:
        inp(M, f"K{r}", round(d["conv"] / d["clics"], 4), PCT2)
    else:
        cell(M, f"K{r}", "–", bg=bg, al="right", color=MUTED)
    for col in "LMN": cell(M, f"{col}{r}", None, bg=bg)
    M.merge_cells(f"L{r}:N{r}")
    cell(M, f"L{r}", uso[k], bg=bg, color=MUTED)
note(M, f"B{b0+5}", f"Fuente: datos reales de Google Ads, Meta Ads y LinkedIn Ads, {PERIODO}. Meta: CPC por clic en el enlace y CR sobre leads calificados.")
note(M, f"B{b0+6}", "Performance Max es una campaña nueva: toma el CPC y la tasa de conversión de las campañas Search B2B de Argentina, con un CTR estimado de 1 %.")

# ---- KPI cards
tc, ta = TOT["Conservador"], TOT["Arriesgado"]
lc, la = LI["Conservador"], LI["Arriesgado"]
card(M, 5, 2, 2, "Flight", f"=TEXT({INI},\"dd/mm\")&\" al \"&TEXT({FIN},\"dd/mm\")", None, L_CYAN, None)
put(M, "B7", f"={DIAS}&\" días · Cyber Monday del \"&TEXT(Supuestos!$C$9,\"dd/mm\")&\" al \"&TEXT(Supuestos!$D$9,\"dd/mm\")", F(size=9, color=MUTED), bg=L_CYAN)
card(M, 5, 3, 4, "Inversión conservador", f"=H{tc}", MONEY, L_VERDE, None)
put(M, "C7", f"=TEXT({MULT['Conservador']},\"0.0\")&\" veces un día común\"", F(size=9, color=MUTED), bg=L_VERDE)
card(M, 5, 6, 7, "Inversión arriesgado", f"=H{ta}", MONEY, L_ROSA, None)
put(M, "F7", f"=TEXT({MULT['Arriesgado']},\"0.0\")&\" veces un día común\"", F(size=9, color=MUTED), bg=L_ROSA)
card(M, 5, 9, 11, "Impresiones", f"=J{tc}", INT, L_VERDE, "Escenario conservador")
card(M, 5, 13, 16, "Alcance LinkedIn", f"=M{lc}", INT, L_CYAN, "Escenario conservador")

# ---- Comparativo de escenarios
band(M, 10, 2, 5, "COMPARATIVO DE ESCENARIOS")
header(M, 11, 2, ["Indicador", "Conservador", "Arriesgado", "Diferencia"])
ind = [("Inversión", "H", MONEY, "t"), ("Impresiones", "J", INT, "t"), ("Clics", "K", INT, "t"),
       ("CTR", "L", PCT2, "t"), ("Alcance LinkedIn", "M", INT, "l"), ("Frecuencia LinkedIn", "N", DEC1, "l"),
       ("Conversiones estimadas", "O", DEC1, "t")]
for i, (n, col, fm, w) in enumerate(ind):
    r = 12 + i
    bg = ZEBRA if i % 2 else None
    cell(M, f"B{r}", n, bg=bg, bold=True)
    cell(M, f"C{r}", f"={col}{tc if w == 't' else lc}", fm, bg)
    cell(M, f"D{r}", f"={col}{ta if w == 't' else la}", fm, bg, bold=True)
    cell(M, f"E{r}", f"=IFERROR(D{r}/C{r}-1,\"\")", PCT0, bg, color=VERDE_INK)
note(M, "B19", "Las conversiones son un indicador secundario: el objetivo de la campaña es presencia como aliado logístico.")

# ---- Inversión por medio
band(M, 10, 8, 12, "INVERSIÓN POR MEDIO")
header(M, 11, 8, ["Medio", "Conservador", "%", "Arriesgado", "%"])
pc_, pa_ = ROWS["Conservador"], ROWS["Arriesgado"]
for i, med in enumerate(["LinkedIn Ads", "Meta Ads", "Google Ads"]):
    r = 12 + i
    bg = ZEBRA if i % 2 else None
    cell(M, f"H{r}", med, bg=bg, bold=True)
    sc = "+".join(f"IF($D${x}=H{r},$H${x},0)" for x in pc_)
    sa = "+".join(f"IF($D${x}=H{r},$H${x},0)" for x in pa_)
    cell(M, f"I{r}", f"={sc}", MONEY, bg)
    cell(M, f"J{r}", f"=IFERROR(I{r}/$I$15,\"\")", PCT0, bg)
    cell(M, f"K{r}", f"={sa}", MONEY, bg)
    cell(M, f"L{r}", f"=IFERROR(K{r}/$K$15,\"\")", PCT0, bg)
total_row(M, 15, 8, 12)
put(M, "H15", "TOTAL", F(bold=True), bg=L_VERDE, box=True)
put(M, "I15", "=SUM(I12:I14)", F(bold=True, color=VERDE_INK), MONEY, L_VERDE, box=True)
put(M, "J15", "=SUM(J12:J14)", F(bold=True), PCT0, L_VERDE, box=True)
put(M, "K15", "=SUM(K12:K14)", F(bold=True, color=VERDE_INK), MONEY, L_VERDE, box=True)
put(M, "L15", "=SUM(L12:L14)", F(bold=True), PCT0, L_VERDE, box=True)
M.conditional_formatting.add("J12:J14", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=VERDE))
M.conditional_formatting.add("L12:L14", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ROSA))
note(M, "H16", "Meta: una sola campaña con presupuesto asignado por grupo de anuncios (público similar y remarketing).")

# ---- Inversión diaria vs. día común
band(M, 21, 2, 9, "INVERSIÓN DIARIA VS. UN DÍA COMÚN DE B2B ARGENTINA")
header(M, 22, 2, ["Escenario", "Día común B2B AR", "Promedio diario", "Diferencia por día", "Veces un día común", "", "Día del Cyber Monday", "Veces un día común"])
for i, escn in enumerate(("Conservador", "Arriesgado")):
    r = 23 + i
    bg = ZEBRA if i % 2 else None
    T = TOT[escn]
    cell(M, f"B{r}", escn, bg=bg, bold=True)
    cell(M, f"C{r}", f"={DIA_COMUN}", MONEY, bg)
    cell(M, f"D{r}", f"=I{T}", MONEY, bg, bold=True)
    cell(M, f"E{r}", f"=D{r}-C{r}", MONEY, bg, color=VERDE_INK)
    cell(M, f"F{r}", f"=IFERROR(D{r}/C{r},\"\")", '0.0"x"', bg, bold=True, al="center")
    cell(M, f"G{r}", None, bg=bg)
    cell(M, f"H{r}", f"=H{T}*MAX(Supuestos!$C${CURVA_PCT}:$L${CURVA_PCT})", MONEY, bg, bold=True)
    cell(M, f"I{r}", f"=IFERROR(H{r}/C{r},\"\")", '0.0"x"', bg, bold=True, al="center")
M.conditional_formatting.add("F23:F24", ColorScaleRule(start_type="min", start_color=L_VERDE, end_type="max", end_color=VERDE))
M.conditional_formatting.add("I23:I24", ColorScaleRule(start_type="min", start_color=L_ROSA, end_type="max", end_color=ROSA))
put(M, "B25", f"=\"En el escenario conservador se invierten \"&TEXT(E23,\"$ #,##0\")&\" más por día que en un día común (\"&TEXT(F23,\"0.0\")&\" veces) y en los días del Cyber Monday \"&TEXT(I23,\"0.0\")&\" veces. En el arriesgado, \"&TEXT(F24,\"0.0\")&\" y \"&TEXT(I24,\"0.0\")&\" veces.\"",
    F(size=9, italic=True, color=MUTED))
note(M, "B26", f"Día común: inversión diaria real de las campañas B2B de Argentina en Google Ads, Meta Ads y LinkedIn Ads, {PERIODO}. En la semana del Cyber Monday todos los anunciantes suben su inversión; el plan acompaña esa suba para no perder visibilidad.")

M.freeze_panes = "C1"

# =====================================================================
# BGT DIARIO (plan por escenario + control de consumo de Milly)
# =====================================================================
B = wb.create_sheet("BGT Diario", 1)
sheet(B, {"A": 2.6, "B": 52, "C": 13, **{L(4 + i): 11 for i in range(22)}})
brand(B, "Inversión diaria  |  Cyber B2B Argentina")
r = 5
plan_bgt = {}
daily = []


def nombre(mr, escn):
    """Campaña, o campaña · grupo de anuncios para Meta, tal como figura en el MP."""
    h = META_HDR[escn]
    return (f"=IF('MP Cyber B2B'!B{mr}=\"\",'MP Cyber B2B'!B{h}&\" · \"&TRIM(SUBSTITUTE('MP Cyber B2B'!C{mr},\"↳\",\"\")),"
            f"'MP Cyber B2B'!B{mr})")


for escn in ("Conservador", "Arriesgado"):
    band(B, r, 2, 13, f"PLAN · ESCENARIO {escn.upper()}"); r += 1
    header(B, r, 2, ["Campaña / grupo de anuncios", "Inversión"] + [None] * 10)
    for i in range(10):
        put(B, f"{L(4+i)}{r}", f"=Supuestos!{L(3+i)}{CURVA_FECHA}", F(bold=True), "ddd dd/mm", VERDE, al="center", box=True)
    r += 1
    b0 = r
    for k, mr in enumerate(plan_rows[escn]):
        bg = ZEBRA if k % 2 else None
        cell(B, f"B{r}", nombre(mr, escn), bg=bg, bold=True)
        cell(B, f"C{r}", f"='MP Cyber B2B'!H{mr}", MONEY, bg)
        for i in range(10):
            cell(B, f"{L(4+i)}{r}", f"=$C{r}*Supuestos!{L(3+i)}${CURVA_PCT}", MONEY, bg)
        r += 1
    daily.append(f"D{b0}:M{r-1}")
    plan_bgt[escn] = (b0, r - 1, r)
    total_row(B, r, 2, 13)
    put(B, f"B{r}", f"Total {escn.lower()}", F(bold=True), bg=L_VERDE, box=True)
    for col in ["C"] + [L(4 + i) for i in range(10)]:
        put(B, f"{col}{r}", f"=SUM({col}{b0}:{col}{r-1})", F(bold=True, color=VERDE_INK if col == "C" else NEGRO), MONEY, L_VERDE, box=True)
    r += 3
for rng in daily:
    B.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=VERDE))

# Control de consumo (estructura de Milly: por día, consumo real y % contra el plan del día)
cs = r
band(B, cs, 2, 25, "CONTROL DE CONSUMO"); cs += 1
put(B, f"B{cs}", f"=\"Escenario controlado: \"&{ESC}", F(size=9, italic=True, color=MUTED)); cs += 1
hdr = ["Campaña / grupo de anuncios", "Inversión"]
for i in range(10):
    hdr += [None, "%"]
header(B, cs, 2, hdr + ["Consumido", "% del total"])
for i in range(10):
    put(B, f"{L(4+2*i)}{cs}", f"=Supuestos!{L(3+i)}{CURVA_FECHA}", F(bold=True), "ddd dd/mm", VERDE, al="center", box=True)
cs += 1
c0 = cs
cb0, cb1, _ = plan_bgt["Conservador"]
ab0, ab1, _ = plan_bgt["Arriesgado"]
n_rows = cb1 - cb0 + 1
for k in range(n_rows):
    rc, ra = cb0 + k, ab0 + k
    bg = ZEBRA if k % 2 else None
    cell(B, f"B{cs}", f"=B{rc}", bg=bg, bold=True)
    cell(B, f"C{cs}", f"=IF({ESC}=\"Arriesgado\",C{ra},C{rc})", MONEY, bg)
    for i in range(10):
        real_col, pct_col, plan_col = L(4 + 2 * i), L(5 + 2 * i), L(4 + i)
        inp(B, f"{real_col}{cs}", None, MONEY)
        cell(B, f"{pct_col}{cs}",
             f"=IF({real_col}{cs}=\"\",\"\",IFERROR({real_col}{cs}/IF({ESC}=\"Arriesgado\",{plan_col}{ra},{plan_col}{rc}),\"\"))",
             PCT0, bg)
    cons_cols = ",".join(f"{L(4+2*i)}{cs}" for i in range(10))
    cell(B, f"X{cs}", f"=SUM({cons_cols})", MONEY, bg, bold=True)
    cell(B, f"Y{cs}", f"=IF(X{cs}=0,\"\",IFERROR(X{cs}/C{cs},\"\"))", PCT0, bg, bold=True)
    cs += 1
total_row(B, cs, 2, 25)
put(B, f"B{cs}", "Total", F(bold=True), bg=L_VERDE, box=True)
put(B, f"C{cs}", f"=SUM(C{c0}:C{cs-1})", F(bold=True, color=VERDE_INK), MONEY, L_VERDE, box=True)
for i in range(10):
    real_col, pct_col, plan_col = L(4 + 2 * i), L(5 + 2 * i), L(4 + i)
    put(B, f"{real_col}{cs}", f"=IF(COUNT({real_col}{c0}:{real_col}{cs-1})=0,\"\",SUM({real_col}{c0}:{real_col}{cs-1}))", F(bold=True), MONEY, L_VERDE, box=True)
    put(B, f"{pct_col}{cs}",
        f"=IF({real_col}{cs}=\"\",\"\",IFERROR({real_col}{cs}/IF({ESC}=\"Arriesgado\",{plan_col}{ab1+1},{plan_col}{cb1+1}),\"\"))",
        F(bold=True), PCT0, L_VERDE, box=True)
put(B, f"X{cs}", f"=SUM(X{c0}:X{cs-1})", F(bold=True, color=VERDE_INK), MONEY, L_VERDE, box=True)
put(B, f"Y{cs}", f"=IF(X{cs}=0,\"\",IFERROR(X{cs}/C{cs},\"\"))", F(bold=True), PCT0, L_VERDE, box=True)
pct_ranges = [f"{L(5+2*i)}{c0}:{L(5+2*i)}{cs}" for i in range(10)] + [f"Y{c0}:Y{cs}"]
for rg in pct_ranges:
    B.conditional_formatting.add(rg, CellIsRule(operator="between", formula=["0.9", "1.1"], font=Font(color=VERDE_INK, bold=True)))
    B.conditional_formatting.add(rg, CellIsRule(operator="notBetween", formula=["0.9", "1.1"], font=Font(color=ROSA, bold=True)))
note(B, f"B{cs+1}", "Cargar el consumo real de cada día. El % compara contra el plan del día: en verde entre 90 % y 110 %, en rosa fuera de ese rango.")
B.freeze_panes = "C1"

# =====================================================================
# AUDIENCIAS
# =====================================================================
A = wb.create_sheet("Audiencias", 2)
sheet(A, {"A": 2.6, "B": 34, "C": 50, "D": 18})
brand(A, "Audiencias y señales")
r = 5
for titulo, filas in (("LINKEDIN ADS · BOOST DE POSTS", ["Audiencia guardada", "Tamaño estimado", "Ubicación", "Cargos / seniority", "Industrias", "Tamaño de empresa"]),
                      ("META ADS · TRÁFICO WEB", ["Grupo 1: público similar", "Base de origen", "Porcentaje de similitud", "Grupo 2: remarketing", "Base de datos", "Tamaño de la base"]),
                      ("GOOGLE ADS · PERFORMANCE MAX", ["Señal: términos de búsqueda", "Señal: remarketing", "Mensaje central", "URL final"])):
    band(A, r, 2, 4, titulo); r += 1
    header(A, r, 2, ["Dato", "Detalle", "Estado"]); r += 1
    for k, f in enumerate(filas):
        bg = ZEBRA if k % 2 else None
        cell(A, f"B{r}", f, bg=bg, bold=True)
        cell(A, f"C{r}", "¿Tu logística aguanta el pico de ventas?" if f == "Mensaje central" else None, bg=bg)
        cell(A, f"D{r}", None, bg=bg)
        r += 1
    r += 1

# Controles internos (en Supuestos, no en el MP)
band(S, 47, 2, 5, "CONTROLES")
ctl = [("Total conservador = presupuesto", f"=IF(ROUND('MP Cyber B2B'!H{tc}-{PRES['Conservador']},0)=0,\"OK\",\"REVISAR\")"),
       ("Total arriesgado = presupuesto", f"=IF(ROUND('MP Cyber B2B'!H{ta}-{PRES['Arriesgado']},0)=0,\"OK\",\"REVISAR\")"),
       ("% de cada escenario = 100 %", f"=IF(AND(ROUND('MP Cyber B2B'!G{tc},4)=1,ROUND('MP Cyber B2B'!G{ta},4)=1),\"OK\",\"REVISAR\")"),
       ("Curva diaria = 100 %", "=IF(ROUND(C41,4)=1,\"OK\",\"REVISAR\")"),
       ("Frecuencia LinkedIn ≤ 10", f"=IF(MAX('MP Cyber B2B'!N{lc},'MP Cyber B2B'!N{la})<=10,\"OK\",\"REVISAR\")")]
for i, (n, f) in enumerate(ctl):
    rr = 48 + i
    cell(S, f"B{rr}", n, bold=True)
    cell(S, f"D{rr}", f, bold=True, al="center")
S.conditional_formatting.add("D48:D52", CellIsRule(operator="equal", formula=['"OK"'], fill=fill(L_VERDE), font=Font(color=VERDE_INK, bold=True)))
S.conditional_formatting.add("D48:D52", CellIsRule(operator="equal", formula=['"REVISAR"'], fill=fill(ROSA), font=Font(color="FFFFFF", bold=True)))

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print("ok", OUT)

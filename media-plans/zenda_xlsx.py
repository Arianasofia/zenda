"""Estilos y helpers de planillas con estética Zenda (openpyxl)."""
import sys, datetime as dt
from openpyxl import Workbook  # noqa: F401
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, ColorScaleRule
from openpyxl.utils import get_column_letter as L


# ---- paleta Zenda
NEGRO, MUTED, VERDE, VERDE_INK, CYAN, ROSA = "0A0A0B", "6B6B6E", "59D7A2", "2FA37A", "95D6EA", "E71CA2"
L_VERDE, L_ROSA, L_CYAN, ZEBRA, GRID, INPUT = "EAF9F3", "FDEAF5", "EAF7FB", "F7F7F8", "E3E3E5", "1155CC"
F = lambda **k: Font(name="Calibri", **{"size": 10, "color": NEGRO, **k})
fill = lambda c: PatternFill("solid", fgColor=c)
side = Side(style="thin", color=GRID)
BOX = Border(left=side, right=side, top=side, bottom=side)
MONEY, INT, DEC1, PCT2, PCT0, DATE = '"$" #,##0', "#,##0", "#,##0.0", "0.00%", "0%", "dd/mm/yyyy"
NUMERIC = {MONEY, INT, DEC1, PCT2, PCT0, DATE, "ddd dd/mm"}



def sheet(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERDE
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def put(ws, ref, v, font=None, fmt=None, bg=None, al=None, box=False):
    c = ws[ref]
    c.value = v
    c.font = font or F()
    if fmt: c.number_format = fmt
    if bg: c.fill = fill(bg)
    if al is None:
        al = "right" if fmt in NUMERIC else "left"
    c.alignment = Alignment(horizontal=al, vertical="center", indent=1 if al == "left" else 0)
    if box: c.border = BOX
    return c


def cell(ws, ref, v, fmt=None, bg=None, bold=False, color=NEGRO, al=None):
    """Celda de tabla: con borde fino."""
    return put(ws, ref, v, F(bold=bold, color=color), fmt, bg, al, box=True)


def inp(ws, ref, v, fmt, note=None):
    """Dato cargado a mano: azul sobre celeste claro."""
    c = put(ws, ref, v, F(color=INPUT), fmt, L_CYAN, box=True)
    if note: c.comment = Comment(note, "Zenda")
    return c


def pend(c, note=None):
    c.fill = fill(L_ROSA)
    c.font = F(color=ROSA, bold=True)
    if note: c.comment = Comment(note, "Zenda")


def band(ws, row, c1, c2, text, size=12, bg=NEGRO, color="FFFFFF", h=24):
    for c in range(c1, c2 + 1):
        ws.cell(row, c).fill = fill(bg)
    put(ws, f"{L(c1)}{row}", text, F(size=size, bold=True, color=color), bg=bg)
    ws.row_dimensions[row].height = h


def header(ws, row, c1, names, bg=VERDE):
    for i, n in enumerate(names):
        c = put(ws, f"{L(c1+i)}{row}", n, F(bold=True), bg=bg, al="center", box=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def groups(ws, row, spans):
    """Fila de grupos de columnas sobre el encabezado (Inversión / Benchmarks / Resultados)."""
    for c1, c2, text, bg, color in spans:
        for c in range(c1, c2 + 1):
            ws.cell(row, c).fill = fill(bg)
            ws.cell(row, c).border = BOX
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        put(ws, f"{L(c1)}{row}", text, F(size=9, bold=True, color=color), bg=bg, al="center", box=True)
    ws.row_dimensions[row].height = 18


def total_row(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        x = ws.cell(row, c)
        x.fill = fill(L_VERDE)
        x.border = Border(left=side, right=side, top=Side(style="medium", color=VERDE), bottom=side)
    ws.row_dimensions[row].height = 20


def brand(ws, title, subtitle=None):
    put(ws, "B2", "zenda.", F(size=12, bold=True, color=VERDE_INK))
    put(ws, "B3", title, F(size=18, bold=True))
    ws.row_dimensions[3].height = 28
    if subtitle:
        put(ws, "B4", subtitle, F(size=9, italic=True, color=MUTED))


def card(ws, row, c1, c2, label, formula, fmt, bg, sub=None):
    """KPI card: label chico gris, número grande, aclaración abajo."""
    for r in range(row, row + 3):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill(bg)
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    put(ws, f"{L(c1)}{row}", label.upper(), F(size=9, bold=True, color=MUTED), bg=bg, al="left")
    v = put(ws, f"{L(c1)}{row+1}", formula, F(size=20, bold=True, color=NEGRO), fmt, bg, al="left")
    if sub:
        put(ws, f"{L(c1)}{row+2}", sub, F(size=9, color=MUTED), bg=bg, al="left")
    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row + 1].height = 32
    ws.row_dimensions[row + 2].height = 16
    return v


def note(ws, ref, text):
    put(ws, ref, text, F(size=9, italic=True, color=MUTED))


